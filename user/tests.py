from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from catalog.models import Category, Product
from core.models import Preorder

from .models import CustomerProfile


User = get_user_model()


class AccountTests(TestCase):
    def test_english_order_pages_translate_account_status_and_catalog_items(self):
        user = User.objects.create_user(
            username="english_customer",
            email="english-customer@example.com",
            password=None,
        )
        category = Category.objects.create(
            name="Подарочные наборы",
            name_en="Gift sets",
        )
        product = Product.objects.create(
            name="Подарочный набор",
            name_en="English gift set",
            category=category,
        )
        order = Preorder.objects.create(
            user=user,
            name="Anna",
            email=user.email,
            phone="+7 900 000-00-00",
            configuration={
                "items": [
                    {
                        "product_id": product.pk,
                        "name": product.name,
                        "quantity": 2,
                        "size": "8",
                        "coating": "painting",
                        "dough": "classic",
                        "packaging": "gift_box",
                        "unit_price": "100.00",
                    }
                ],
                "delivery": "pickup",
                "payment": "card",
            },
            unit_price="100.00",
            discount_percent=0,
            delivery_cost="0.00",
            total_price="200.00",
            status=Preorder.Status.PRODUCTION,
            consent_given=True,
        )
        self.client.force_login(user)
        self.client.post(
            reverse("set_language"),
            {"language": "en", "next": reverse("accounts:orders")},
        )

        orders_page = self.client.get(reverse("accounts:orders"))
        self.assertContains(orders_page, "Your orders")
        self.assertContains(orders_page, "1 order")
        self.assertContains(orders_page, "English gift set × 2")
        self.assertContains(orders_page, "In production")

        detail_page = self.client.get(
            reverse("accounts:order_detail", args=[order.pk])
        )
        self.assertContains(detail_page, "Order progress")
        self.assertContains(detail_page, "Hand-painted icing")
        self.assertContains(detail_page, "Pickup")

    def test_orders_page_shows_only_orders_owned_by_the_logged_in_user(self):
        user = User.objects.create_user(
            username="order_customer",
            email="order-customer@example.com",
            password="A-long-password-482!",
        )
        other_user = User.objects.create_user(
            username="other_customer",
            email="other-customer@example.com",
            password="A-long-password-583!",
        )
        order_values = {
            "name": "Анна",
            "email": "order-customer@example.com",
            "phone": "+7 900 000-00-00",
            "configuration": {
                "items": [{"name": "Подарочный набор", "quantity": 2}],
                "delivery": "pickup",
                "payment": "card",
            },
            "unit_price": "1490.00",
            "discount_percent": 0,
            "delivery_cost": "0.00",
            "total_price": "2980.00",
            "consent_given": True,
        }
        order = Preorder.objects.create(user=user, **order_values)
        Preorder.objects.create(
            user=other_user,
            **{**order_values, "email": "other-customer@example.com"},
        )
        self.client.force_login(user)

        response = self.client.get(reverse("accounts:orders"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, f"Заказ #{order.pk}")
        self.assertContains(response, "Подарочный набор × 2")
        self.assertContains(response, "1 заказ")
        self.assertNotContains(response, "Заказ #2")

    def test_dashboard_shows_current_user_order_count_and_recent_order(self):
        user = User.objects.create_user(
            username="dashboard_customer",
            email="dashboard-customer@example.com",
            password="A-long-password-482!",
        )
        order = Preorder.objects.create(
            user=user,
            name="Анна",
            email=user.email,
            phone="+7 900 000-00-00",
            configuration={"quantity": 4, "size": "8"},
            unit_price="100.00",
            discount_percent=0,
            delivery_cost="700.00",
            total_price="1100.00",
            consent_given=True,
        )
        self.client.force_login(user)

        response = self.client.get(reverse("accounts:dashboard"))

        self.assertContains(response, "Мои заказы")
        self.assertContains(response, "1 заказ")
        self.assertContains(response, f"Заказ #{order.pk}")
        self.assertContains(response, "4 шт. · 8 см")

    def test_order_detail_shows_items_and_production_tracking(self):
        user = User.objects.create_user(
            username="tracking_customer",
            email="tracking-customer@example.com",
            password="A-long-password-482!",
        )
        order = Preorder.objects.create(
            user=user,
            name="Анна",
            email=user.email,
            phone="+7 900 000-00-00",
            address="Москва, ул. Примерная, 1",
            configuration={
                "items": [
                    {
                        "name": "Подарочный набор",
                        "quantity": 2,
                        "size": "8",
                        "unit_price": "1490.00",
                    },
                ],
                "delivery": "moscow",
                "payment": "card",
            },
            unit_price="1490.00",
            discount_percent=0,
            delivery_cost="700.00",
            total_price="3680.00",
            status=Preorder.Status.PRODUCTION,
            consent_given=True,
        )
        self.client.force_login(user)

        response = self.client.get(
            reverse("accounts:order_detail", args=[order.pk])
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Состав заказа")
        self.assertContains(response, "Подарочный набор")
        self.assertContains(response, "В производстве")
        self.assertContains(response, "Доставка по Москве")
        self.assertContains(response, 'class="account-order-step is-complete"')
        self.assertContains(response, 'class="account-order-step is-current"')

    def test_customer_cannot_view_another_users_order_detail(self):
        user = User.objects.create_user(
            username="order_owner",
            email="order-owner@example.com",
            password="A-long-password-482!",
        )
        other_user = User.objects.create_user(
            username="order_viewer",
            email="order-viewer@example.com",
            password="A-long-password-583!",
        )
        order = Preorder.objects.create(
            user=user,
            name="Анна",
            email=user.email,
            phone="+7 900 000-00-00",
            configuration={"quantity": 1},
            unit_price="100.00",
            discount_percent=0,
            delivery_cost="700.00",
            total_price="800.00",
            consent_given=True,
        )
        self.client.force_login(other_user)

        response = self.client.get(
            reverse("accounts:order_detail", args=[order.pk])
        )

        self.assertEqual(response.status_code, 404)

    def test_customer_can_save_delivery_address_for_future_orders(self):
        user = User.objects.create_user(
            username="address_customer",
            email="address-customer@example.com",
            password="A-long-password-482!",
        )
        self.client.force_login(user)

        response = self.client.post(
            reverse("accounts:addresses"),
            {"delivery_address": "Москва, ул. Примерная, 12, кв. 4"},
        )

        self.assertRedirects(response, reverse("accounts:addresses"))
        self.assertEqual(
            user.customer_profile.delivery_address,
            "Москва, ул. Примерная, 12, кв. 4",
        )
        self.assertContains(
            self.client.get(reverse("accounts:addresses")),
            "Москва, ул. Примерная, 12, кв. 4",
        )

    def test_registration_shows_customer_type_and_hides_company_fields_initially(self):
        response = self.client.get("/account/register/")

        self.assertContains(response, "Физическое лицо")
        self.assertContains(response, "Юридическое лицо")
        self.assertContains(response, 'id="registration-company-fields" hidden')

    def test_registration_creates_user_and_logs_them_in(self):
        response = self.client.post(
            "/account/register/",
            {
                "username": "new_customer",
                "first_name": "Анна",
                "last_name": "Иванова",
                "email": "anna@example.com",
                "password1": "A-long-password-482!",
                "password2": "A-long-password-482!",
                "customer_type": CustomerProfile.CustomerType.INDIVIDUAL,
            },
        )

        self.assertRedirects(response, "/account/")
        user = User.objects.get(username="new_customer")
        self.assertEqual(user.email, "anna@example.com")
        self.assertEqual(
            user.customer_profile.customer_type,
            CustomerProfile.CustomerType.INDIVIDUAL,
        )
        self.assertEqual(int(self.client.session["_auth_user_id"]), user.pk)

    def test_legal_entity_registration_requires_all_company_fields(self):
        response = self.client.post(
            "/account/register/",
            {
                "username": "legal_customer",
                "email": "legal@example.com",
                "password1": "A-long-password-482!",
                "password2": "A-long-password-482!",
                "customer_type": CustomerProfile.CustomerType.LEGAL_ENTITY,
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertFalse(User.objects.filter(username="legal_customer").exists())
        for field in ("company_name", "inn", "kpp", "legal_address", "contact_name", "phone"):
            with self.subTest(field=field):
                self.assertIn("Обязательное поле для юридического лица.", response.context["form"][field].errors)

    def test_legal_entity_registration_saves_company_details(self):
        response = self.client.post(
            "/account/register/",
            {
                "username": "legal_customer",
                "email": "legal@example.com",
                "password1": "A-long-password-482!",
                "password2": "A-long-password-482!",
                "customer_type": CustomerProfile.CustomerType.LEGAL_ENTITY,
                "company_name": "ООО Пример",
                "inn": "1234567890",
                "kpp": "123456789",
                "legal_address": "Москва, ул. Примерная, д. 1",
                "contact_name": "Анна Иванова",
                "phone": "+7 900 123-45-67",
            },
        )

        self.assertRedirects(response, "/account/")
        profile = CustomerProfile.objects.get(user__username="legal_customer")
        self.assertEqual(profile.customer_type, CustomerProfile.CustomerType.LEGAL_ENTITY)
        self.assertEqual(profile.company_name, "ООО Пример")
        self.assertEqual(profile.inn, "1234567890")
        self.assertEqual(profile.kpp, "123456789")
        self.assertEqual(profile.legal_address, "Москва, ул. Примерная, д. 1")
        self.assertEqual(profile.contact_name, "Анна Иванова")
        self.assertEqual(profile.phone, "+7 900 123-45-67")

    def test_legal_entity_registration_validates_tax_identifier_formats(self):
        response = self.client.post(
            "/account/register/",
            {
                "username": "invalid_legal_customer",
                "email": "invalid-legal@example.com",
                "password1": "A-long-password-482!",
                "password2": "A-long-password-482!",
                "customer_type": CustomerProfile.CustomerType.LEGAL_ENTITY,
                "company_name": "ООО Пример",
                "inn": "123",
                "kpp": "12",
                "legal_address": "Москва",
                "contact_name": "Анна Иванова",
                "phone": "+7 900 123-45-67",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertFalse(User.objects.filter(username="invalid_legal_customer").exists())
        self.assertContains(response, "Введите 10 или 12 цифр.")
        self.assertContains(response, "Введите 9 цифр.")

    def test_admin_can_export_selected_legal_entities_as_xlsx(self):
        admin_user = User.objects.create_superuser(
            username="admin",
            email="admin@example.com",
            password="A-long-password-482!",
        )
        legal_user = User.objects.create_user(
            username="legal_export",
            email="legal-export@example.com",
        )
        CustomerProfile.objects.create(
            user=legal_user,
            customer_type=CustomerProfile.CustomerType.LEGAL_ENTITY,
            company_name="=ООО Опасная формула",
            inn="1234567890",
            kpp="123456789",
            legal_address="Москва",
            contact_name="Иван Иванов",
            phone="+7 900 123-45-67",
        )
        individual_user = User.objects.create_user(username="individual_export")
        individual_profile = CustomerProfile.objects.create(user=individual_user)
        self.client.force_login(admin_user)

        response = self.client.post(
            reverse("admin:user_customerprofile_changelist"),
            {
                "action": "export_legal_entities_xlsx",
                "_selected_action": [str(legal_user.customer_profile.pk), str(individual_profile.pk)],
                "index": "0",
            },
        )

        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(rows[0][0:6], ("Контактное лицо", "Телефон", "Организация", "ИНН", "КПП", "Юридический адрес"))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][2], "'=ООО Опасная формула")

    def test_non_staff_users_cannot_access_customer_admin(self):
        response = self.client.get(reverse("admin:user_customerprofile_changelist"))

        self.assertEqual(response.status_code, 302)
        self.assertIn("/admin/login/", response["Location"])

    def test_registration_rejects_duplicate_email(self):
        User.objects.create_user(
            username="first_customer",
            email="taken@example.com",
            password="A-long-password-482!",
        )

        response = self.client.post(
            "/account/register/",
            {
                "username": "second_customer",
                "first_name": "",
                "last_name": "",
                "email": "TAKEN@example.com",
                "password1": "Another-long-password-583!",
                "password2": "Another-long-password-583!",
                "customer_type": CustomerProfile.CustomerType.INDIVIDUAL,
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Аккаунт с такой почтой уже существует.")
        self.assertFalse(User.objects.filter(username="second_customer").exists())

    def test_login_logout_and_protected_dashboard(self):
        user = User.objects.create_user(
            username="returning_customer",
            password="A-long-password-482!",
        )

        response = self.client.get("/account/")
        self.assertRedirects(response, "/account/login/?next=/account/")

        response = self.client.post(
            "/account/login/",
            {"username": user.username, "password": "A-long-password-482!"},
        )
        self.assertRedirects(response, "/account/")
        self.assertContains(self.client.get("/account/"), "Личный кабинет")

        response = self.client.post("/account/logout/")
        self.assertRedirects(response, "/")
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_dashboard_saves_profile_details(self):
        user = User.objects.create_user(
            username="profile_customer",
            email="before@example.com",
            password="A-long-password-482!",
        )
        self.client.force_login(user)

        response = self.client.post(
            "/account/profile/",
            {
                "profile-first_name": "Мария",
                "profile-last_name": "Петрова",
                "profile-email": "maria@example.com",
            },
        )

        self.assertRedirects(response, "/account/profile/")
        user.refresh_from_db()
        self.assertEqual(user.first_name, "Мария")
        self.assertEqual(user.last_name, "Петрова")
        self.assertEqual(user.email, "maria@example.com")

    def test_account_sections_have_distinct_pages_and_active_navigation(self):
        user = User.objects.create_user(
            username="section_customer",
            password="A-long-password-482!",
        )
        self.client.force_login(user)

        pages = (
            ("/account/", "Личный кабинет", "/account/orders/"),
            ("/account/orders/", "Мои заказы", "/account/favorites/"),
            ("/account/favorites/", "Избранное", "/account/addresses/"),
            ("/account/addresses/", "Адрес доставки", "/account/profile/"),
            ("/account/profile/", "Мои данные", "/account/"),
        )
        for url, heading, next_page in pages:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, f"<h1>{heading}</h1>", html=True)
                self.assertContains(response, f'href="{next_page}"')

    def test_account_sections_require_authentication(self):
        for url in (
            "/account/orders/",
            "/account/addresses/",
            "/account/profile/",
        ):
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertRedirects(response, f"/account/login/?next={url}")
